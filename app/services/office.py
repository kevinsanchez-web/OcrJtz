"""Documentos de Office (.docx y .xlsx): su texto y sus tablas se leen directamente, sin OCR.

Un .docx/.xlsx es un ZIP con XML dentro, así que no hay nada que "reconocer": se extrae lo escrito. Se limitan
el tamaño descomprimido y el número de filas para que un archivo malicioso (zip bomb) no agote la memoria.
"""
import io
import logging
import zipfile
from typing import List, Optional, Tuple

from app.schemas import Table

logger = logging.getLogger("document-api")

MAX_UNCOMPRESSED = 150 * 1024 * 1024
MAX_ROWS, MAX_COLS, MAX_SHEETS = 5000, 60, 20
LEGACY_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # .doc / .xls antiguos (OLE2)


def office_kind(contents: bytes) -> Optional[str]:
    """'docx', 'xlsx', 'legacy' (.doc/.xls antiguos) o None si no es un archivo de Office."""
    if contents[:8] == LEGACY_MAGIC:
        return "legacy"
    if contents[:2] != b"PK":
        return None
    try:
        with zipfile.ZipFile(io.BytesIO(contents)) as z:
            infos = z.infolist()
            if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED or len(infos) > 5000:
                return None
            names = {i.filename for i in infos}
    except Exception:
        return None
    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    return None


def _clean(v) -> str:
    return "" if v is None else str(v).replace("\n", " ").strip()


def read_docx(contents: bytes) -> List[Tuple[str, List[Table]]]:
    from docx import Document
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(contents))
    lines: List[str] = []
    tables: List[Table] = []
    for child in doc.element.body.iterchildren():  # párrafos y tablas en el orden del documento
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, doc).text.strip()
            if text:
                lines.append(text)
        elif tag == "tbl":
            rows = [[_clean(c.text) for c in r.cells] for r in DocxTable(child, doc).rows]
            rows = [r for r in rows if any(r)]
            if rows:
                tables.append(Table(rows=rows, source="grid", box=[0.0, 0.0, 1.0, 1.0]))
                lines.extend(" | ".join(r) for r in rows)
    for sec in doc.sections:  # encabezados y pies suelen traer datos (nombre de la empresa, NIT...)
        for part in (sec.header, sec.footer):
            lines.extend(p.text.strip() for p in part.paragraphs if p.text.strip())
    return [("\n".join(lines), tables)]


def read_xlsx(contents: bytes) -> List[Tuple[str, List[Table]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    pages: List[Tuple[str, List[Table]]] = []
    try:
        for ws in list(wb.worksheets)[:MAX_SHEETS]:
            rows: List[List[str]] = []
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_ROWS:
                    break
                cells = [_clean(c) for c in row[:MAX_COLS]]
                if any(cells):
                    while cells and not cells[-1]:
                        cells.pop()
                    rows.append(cells)
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            text = f"Hoja: {ws.title}\n" + "\n".join(" | ".join(r) for r in rows)
            pages.append((text, [Table(rows=rows, source="grid", box=[0.0, 0.0, 1.0, 1.0])] if width >= 2 else []))
    finally:
        wb.close()
    return pages


def read_office(contents: bytes, kind: str) -> List[Tuple[str, List[Table]]]:
    """[(texto, tablas)] por página (docx: una; xlsx: una por hoja)."""
    return read_docx(contents) if kind == "docx" else read_xlsx(contents)
